# -*- coding: utf-8 -*-
r"""
Daily Container Stock vs Booking report  (all 9 types, all locations)
====================================================================
Inputs : STAYING/stock .xls  +  BKG+PD booking .xls
Outputs: Stock_Daily_<M-D>.xlsx   (Control / StockRaw / BookingRaw / Summary, formula-driven)
         Stock_Daily_<M-D>.html   (self-contained dashboard, same numbers)
         Stock_Daily_<M-D>.model.json

Run (explicit files):  python build_stock_daily.py "9-10-STAYING.xls" "9-10-BKG+PD.xls" [YYYY-MM-DD]
Run (auto from input/ folder, newest STAYING/BKG pair):  python build_stock_daily.py [YYYY-MM-DD]
Then: powershell -File recalc_xlsx.ps1 -Path .\Stock_Daily_<M-D>.xlsx
"""
from __future__ import annotations

import datetime as dt
import html
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FOLDER = Path(__file__).resolve().parent
INPUT_DIR = FOLDER / "input"          # daily STAYING/booking exports get moved here

BKG_TYPE_ORDER = ["GP22", "GP42", "GP45", "RE22", "RE45", "UT22", "UT42", "PC22", "PC42"]
CODE2DISP = {"GP22": "20'GP", "GP42": "40'GP", "GP45": "40'HC", "RE22": "20'RE", "RE45": "40'RH",
             "UT22": "20'OT", "UT42": "40'OT", "PC22": "20'FR", "PC42": "40'FR"}
SIZE2DISP = {"22GP": "20'GP", "42GP": "40'GP", "45GP": "40'HC", "22RE": "20'RE", "45RE": "40'RH",
             "22UT": "20'OT", "42UT": "40'OT", "22PC": "20'FR", "42PC": "40'FR"}
DISP_ORDER = ["20'GP", "40'GP", "40'HC", "20'RE", "40'RH", "20'OT", "40'OT", "20'FR", "40'FR"]
RE_DISP = {"20'RE", "40'RH"}
OTFR_DISP = {"20'OT", "40'OT", "20'FR", "40'FR"}
DISP2SIZE = {v: k for k, v in SIZE2DISP.items()}   # 20'GP -> 22GP, ... (headers of the no-booking table)

# confirmed conventions
MERGES = {"BKK04": "BKK01"}                     # relabel -> merged block
MERGE_LABEL = {"BKK01": "BKK01+BKK04"}
BKK_ZONE_FORCE = {"LCH54", "LCH55"}                      # LCH-named but BKK zone
# these locations always stay in the full block grid even with zero bookings
# (never routed into the compact "no active bookings" section) - user request
ALWAYS_FULL_BLOCK = {"BKK02", "LCH55"}
# fixed lead-in order per zone (any other codes present fall in after these,
# alphabetically) - codes not listed keep their default alphabetical spot.
# BKK25 then BKK27 is deliberate: with 3 lanes, index i and i+3 land in the
# same lane one row-group apart, so BKK27 (index 3) renders directly under
# BKK25 (index 0) - user request.
LOCATION_ORDER = {"BKK": ["BKK25", "BKK01", "BKK02", "BKK27", "LCH54", "LCH55"], "LCH": []}
# yard/depot name shown after the location code in each block's title band
# (keyed by the display code, e.g. the merged "BKK01+BKK04") - user request
YARD_NAME = {
    "BKK01+BKK04": "PAT TERMINAL 1 & 2",
    "BKK02": "UNITHAI CONTAINER TERMINAL",
    "BKK25": "Smart Logistics Service (Thailand ) Co.,Ltd.",
    "BKK27": "B.C. DEPOT CO.,LTD. ( Bang-Na KM.18)",
    "LCH55": "ESCO LKR",
    "LCH27": "HAST Logistics Co.,Ltd.",
    "LCH28": "CELLO",
    "LCHY5": "PW DEPOT CO., LTD.",
}
# location codes to drop from the report entirely (not real/tracked yards) -
# user request
EXCLUDE_LOCATIONS = {"LCH20"}

NAVY = "0C2340"
ZONE_BAND_FILL = "FFFF00"    # zone band rows - user request 2026-09-23
# zone-band layout gaps: extra blank block-slots inserted before the named
# code so it lands directly under an earlier block in the same lane (index i
# and i+3 share a lane, one row-group apart) - user request
LAYOUT_GAPS = {"LCH": {"LCHY5": 1}}
STOCK_BG = "F3F5F8"
AV_BG = "FFF6C8"
RE_FONT = "1F4E9C"
OTFR_FONT = "7A4B12"
NEG_FONT = "C22A2A"

BLOCK_H = 13
LANE_STARTS = {0: 1, 1: 13, 2: 25}             # A / M / Y
GAP_COLS = [12, 24]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


# --------------------------------------------------------------------------- #
def zone_of(code: str) -> str:
    if code in BKK_ZONE_FORCE:
        return "BKK"
    return "LCH" if code.upper().startswith("LCH") else "BKK"


def load_stock(path: Path):
    df = pd.read_excel(path, header=0)
    df = df[df["Location"].notna() & df["Size/Type"].notna()]
    if "Full/Empty" in df.columns:
        df = df[df["Full/Empty"].astype(str).str.strip().str.upper() == "E"]
    df["loc"] = df["Location"].astype(str).str.strip().str.upper().replace(MERGES)
    df["disp"] = df["Size/Type"].astype(str).str.strip().str.upper().map(SIZE2DISP)
    df = df[df["disp"].notna()]
    raw = df[["loc", "disp"]].copy()
    agg = {}
    for (loc, disp), n in df.groupby(["loc", "disp"]).size().items():
        agg[(loc, disp)] = int(n)
    return agg, raw


def load_booking(path: Path) -> pd.DataFrame:
    # the booking workbook's sheet order/count varies by day (sometimes a small
    # pivot/summary sheet comes first) - scan every sheet, at header row 0 then
    # 1, until one actually has the 'Pickup' column.
    df = None
    for sheet in pd.ExcelFile(path).sheet_names:
        for header in (0, 1):
            cand = pd.read_excel(path, sheet_name=sheet, header=header)
            if "Pickup" in cand.columns and "TRAN DT" in cand.columns:
                df = cand
                break
        if df is not None:
            break
    if df is None:
        raise SystemExit(f"booking file {path.name}: no sheet with 'Pickup'+'TRAN DT' columns found")
    # some daily exports repeat the header row mid-table (page-break artifact) -
    # coerce TRAN DT to numeric and drop anything that isn't a real YYYYMMDD value
    tran = pd.to_numeric(df["TRAN DT"], errors="coerce")
    df = df[df["Pickup"].notna() & df["Pickup"].astype(str).str.strip().ne("Pickup")
            & tran.notna()].copy()
    tran = tran.loc[df.index]
    loc = df["Pickup"].astype(str).str.strip().str.upper().replace(MERGES)
    date = pd.to_datetime(tran.astype("int64").astype(str), format="%Y%m%d")
    out = pd.DataFrame({"Location": loc.values, "Date": date.values})
    for code in BKG_TYPE_ORDER:
        if code in df.columns:
            out[CODE2DISP[code]] = pd.to_numeric(df[code], errors="coerce").fillna(0).astype(int).values
        else:
            out[CODE2DISP[code]] = 0
    return out.reset_index(drop=True)


def load_all_bookings(paths: list[Path]) -> pd.DataFrame:
    """Load and concatenate one or more booking files.

    Some days the source system splits the export into a "PENDING" file
    (already-overdue bookings, dated before today) and a forward-looking
    "BKG-3WK" file (today onward) - these are complementary, not duplicates
    (confirmed 2026-09-23: zero row overlap between the two), so all given
    files are simply concatenated.
    """
    parts = [load_booking(p) for p in paths]
    return pd.concat(parts, ignore_index=True) if len(parts) > 1 else parts[0]


SIZE_CODES = list(SIZE2DISP)   # 22GP, 42GP, 45GP, 22RE, 45RE, 22UT, 42UT, 22PC, 42PC


def load_empty_repo(path: Path) -> pd.DataFrame:
    """Inbound empty-repositioning list (the "EMPTY" export: move code VED, discharge at TH).

    One row per container: Location = destination yard, EQ Date = arrival date,
    plus Vessel/Voyage/P.O.L. A trailing summary row (no container no.) is dropped.
    Returns columns Location, ETA, Vessel, Voyage, POL, Type (size code, e.g. 22GP).
    """
    df = pd.read_excel(path)
    df = df[df["Container No."].notna() & df["Type Size"].notna() & df["Location"].notna()].copy()
    df["Type"] = df["Type Size"].astype(str).str.strip().str.upper()
    unknown = sorted(set(df["Type"]) - set(SIZE_CODES))
    if unknown:
        print("WARNING: empty-repo size codes not in report columns, skipped:", unknown)
        df = df[df["Type"].isin(SIZE_CODES)]
    out = pd.DataFrame({
        "Location": df["Location"].astype(str).str.strip().replace(MERGES).replace(MERGE_LABEL),
        "ETA": pd.to_datetime(df["EQ Date"]).dt.normalize(),
        "Vessel": df["Vessel"].astype(str).str.strip(),
        "Voyage": df["Voyage"].astype(str).str.strip(),
        "POL": df["P.O.L"].astype(str).str.strip(),
        "Type": df["Type"],
    })
    return out.reset_index(drop=True)


def empty_repo_groups(er: pd.DataFrame) -> list[dict]:
    """Group the raw rows by ETA / location / vessel / voyage / POL with per-size counts."""
    rows = []
    keys = ["ETA", "Location", "Vessel", "Voyage", "POL"]
    for k, g in er.groupby(keys, sort=True):
        counts = g["Type"].value_counts()
        rows.append({"eta": pd.Timestamp(k[0]).strftime("%Y-%m-%d"), "loc": k[1], "vessel": k[2],
                     "voyage": k[3], "pol": k[4],
                     "types": {c: int(counts.get(c, 0)) for c in SIZE_CODES},
                     "total": int(len(g))})
    return rows


def week_frame(today: dt.date):
    mon1 = today - dt.timedelta(days=today.weekday())
    weeks = [(mon1 + dt.timedelta(days=7 * i), mon1 + dt.timedelta(days=7 * i + 5)) for i in range(4)]
    return weeks, today + dt.timedelta(days=1)


def ordered_locations(stock_agg, bkg):
    codes = ({loc for (loc, _) in stock_agg} | set(bkg["Location"].unique())) - EXCLUDE_LOCATIONS

    def zone_order(zname):
        zcodes = [c for c in codes if zone_of(c) == zname]
        pref = LOCATION_ORDER.get(zname, [])
        head = [c for c in pref if c in zcodes]
        tail = sorted(c for c in zcodes if c not in pref)
        return head + tail

    bkk = zone_order("BKK")
    lch = zone_order("LCH")
    return [(c, "BKK") for c in bkk] + [(c, "LCH") for c in lch]


def build_model(stock_agg, bkg: pd.DataFrame, today: dt.date):
    weeks, tomorrow = week_frame(today)
    sat1 = weeks[0][1]
    b = bkg.copy()
    b["d"] = pd.to_datetime(b["Date"]).dt.date
    locs = []
    for code, zone in ordered_locations(stock_agg, bkg):
        sub = b[b["Location"] == code]
        types = {}
        for disp in DISP_ORDER:
            stock = stock_agg.get((code, disp), 0)
            pend = int(sub.loc[sub["d"] < today, disp].sum()) if len(sub) else 0
            t0 = int(sub.loc[sub["d"] == today, disp].sum()) if len(sub) else 0
            rest = int(sub.loc[(sub["d"] > today) & (sub["d"] <= sat1), disp].sum()) if len(sub) else 0
            wk = [t0 + rest]
            for a, bb in weeks[1:]:
                wk.append(int(sub.loc[(sub["d"] >= a) & (sub["d"] <= bb), disp].sum()) if len(sub) else 0)
            av, run = [], stock - pend
            for i in range(4):
                run -= wk[i]
                av.append(run)
            types[disp] = dict(stock=stock, pending=pend, today=t0, wk1rest=rest, wk=wk, av=av)
        locs.append(dict(code=MERGE_LABEL.get(code, code), key=code, zone=zone, types=types,
                         has_booking=len(sub) > 0))
    return dict(date=today.isoformat(),
                weeks=[[a.isoformat(), b_.isoformat()] for a, b_ in weeks],
                iso_weeks=[a.isocalendar()[1] for a, _ in weeks],
                tomorrow=tomorrow.isoformat(), locations=locs)


# --------------------------------------------------------------------------- #
def mn(ref):
    return f'CHOOSE(MONTH({ref}),"' + '","'.join(MONTHS) + '")'


def range_label(s, e):
    return (f'IF(MONTH({s})=MONTH({e}),DAY({s})&"-"&DAY({e})&" "&{mn(e)},'
            f'DAY({s})&" "&{mn(s)}&"-"&DAY({e})&" "&{mn(e)})')


def build_excel(model, stock_agg, stock_raw, bkg, out: Path, override_date: dt.date | None = None,
                empty_raw: pd.DataFrame | None = None):
    wb = Workbook()
    F = Font(name="Calibri", size=11)
    FB = Font(name="Calibri", size=11, bold=True)
    thin = Side(style="thin", color="C9D2DF")
    med = Side(style="medium", color=NAVY)

    # Control ------------------------------------------------------------
    ws = wb.active
    ws.title = "Control"
    ws["A1"] = "Report date"
    # override_date is used when the report is being labeled for a date other
    # than the machine's live clock (e.g. built ahead using a fixed date) -
    # in that case B1 is a literal date, not a live =TODAY() formula, since a
    # live formula would silently disagree with the intended report date.
    ws["B1"] = dt.datetime.combine(override_date, dt.time()) if override_date else "=TODAY()"
    ws["B1"].number_format = "yyyy-mm-dd"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "WK#", "Monday", "Saturday", "Week no"
    ws["B4"] = "=$B$1-(WEEKDAY($B$1,2)-1)"
    ws["C4"] = "=$B$4+5"
    for r in (5, 6, 7):
        ws[f"B{r}"] = f"=$B${r-1}+7"
        ws[f"C{r}"] = f"=$B${r}+5"
    for r in (4, 5, 6, 7):
        ws[f"A{r}"] = f"=WEEKNUM($B${r},21)"
        ws[f"D{r}"] = f"=WEEKNUM($B${r},21)"
        ws[f"B{r}"].number_format = ws[f"C{r}"].number_format = "yyyy-mm-dd"
    ws["A9"], ws["B9"] = "Tomorrow", "=$B$1+1"
    ws["B9"].number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 16
    for L in ("B", "C", "D"):
        ws.column_dimensions[L].width = 13
    for row in ws.iter_rows():
        for c in row:
            c.font = F

    # StockRaw ---------------------------------------------------------
    sr = wb.create_sheet("StockRaw")
    sr.append(["Location", "Type"])
    for _, rr in stock_raw.iterrows():
        sr.append([rr["loc"], rr["disp"]])
    for row in sr.iter_rows():
        for c in row:
            c.font = F

    # BookingRaw -----------------------------------------------------
    br = wb.create_sheet("BookingRaw")
    br.append(["Location", "Date"] + DISP_ORDER)
    for _, rr in bkg.iterrows():
        br.append([rr["Location"], pd.to_datetime(rr["Date"]).to_pydatetime()]
                  + [int(rr[d]) for d in DISP_ORDER])
    for row in br.iter_rows():
        for c in row:
            c.font = F
        if row[0].row > 1:
            row[1].number_format = "yyyy-mm-dd"
    for i, w in enumerate([16, 12] + [7] * 9):
        br.column_dimensions[get_column_letter(i + 1)].width = w

    # Summary -----------------------------------------------------
    sm = wb.create_sheet("Summary")
    last_col = LANE_STARTS[2] + 10
    for i in range(1, last_col + 1):
        L = get_column_letter(i)
        if i in GAP_COLS:
            sm.column_dimensions[L].width = 2.3
        elif (i - 1) % 12 == 0:
            sm.column_dimensions[L].width = 5.5
        elif (i - 2) % 12 == 0:
            sm.column_dimensions[L].width = 20
        else:
            sm.column_dimensions[L].width = 6.2

    title = ('="Daily Container Stock vs Booking   as of "&DAY(Control!$B$1)&" "&'
             + mn("Control!$B$1") + '&" "&YEAR(Control!$B$1)')
    sm.cell(1, 1, title).font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    sm.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    for cc in range(1, last_col + 1):
        sm.cell(1, cc).fill = PatternFill("solid", fgColor=NAVY)

    mt = {l["code"]: l for l in model["locations"]}
    order = model["locations"]
    with_bkg = [l for l in order if l.get("has_booking", True) or l["key"] in ALWAYS_FULL_BLOCK]
    no_bkg = [l for l in order if not l.get("has_booking", True) and l["key"] not in ALWAYS_FULL_BLOCK]
    zones = []
    for zname, zlabel in (("BKK", "BKK  —  Bangkok area depots"),
                          ("LCH", "LCH  —  Laem Chabang area depots")):
        zlocs = [l for l in with_bkg if l["zone"] == zname]
        if zlocs:
            zones.append((zname, zlabel, zlocs))

    r = 2   # row 2 is blank (the "Report date / Legend" subtitle was dropped - user request)
    if no_bkg:
        r += 1
        band = sm.cell(r, 1, "Stock on hand only  —  no active bookings this period")
        band.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        band.alignment = Alignment(horizontal="left", vertical="center")
        sm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
        for cc in range(1, last_col + 1):
            sm.cell(r, cc).fill = PatternFill("solid", fgColor="6B7280")
        sm.row_dimensions[r].height = 15
        r += 1
        head_row = r
        sm.cell(head_row, 1, "Zone").font = FB
        sm.cell(head_row, 2, "Location").font = FB
        for i, disp in enumerate(DISP_ORDER):
            hc = sm.cell(head_row, 3 + i, DISP2SIZE[disp])
            col = RE_FONT if disp in RE_DISP else OTFR_FONT if disp in OTFR_DISP else NAVY
            hc.font = Font(name="Calibri", size=11, bold=True, color=col)
            hc.alignment = Alignment(horizontal="center")
        # Total sits in L:M merged - column L is the 2.3-wide lane gap, too narrow on its own
        tc = sm.cell(head_row, 12, "Total")
        tc.font = FB
        tc.alignment = Alignment(horizontal="center")
        sm.merge_cells(start_row=head_row, start_column=12, end_row=head_row, end_column=13)
        r += 1
        top_row = r
        for loc in no_bkg:
            sm.cell(r, 1, loc["zone"]).font = F
            sm.cell(r, 2, loc["code"]).font = FB
            for i, disp in enumerate(DISP_ORDER):
                cell = sm.cell(r, 3 + i, loc["types"][disp]["stock"])
                cell.font = F
                cell.alignment = Alignment(horizontal="center")
            for cc in range(1, last_col + 1):
                sm.cell(r, cc).fill = PatternFill("solid", fgColor=STOCK_BG)
            tot = sm.cell(r, 12, f"=SUM(C{r}:K{r})")
            tot.font = FB
            tot.alignment = Alignment(horizontal="center")
            sm.merge_cells(start_row=r, start_column=12, end_row=r, end_column=13)
            r += 1
        bot_row = r - 1
        for rr in range(head_row, bot_row + 1):
            for cc in range(1, 13 + 1):
                cell = sm.cell(rr, cc)
                cell.border = Border(
                    left=med if cc == 1 else thin,
                    right=med if cc == 13 else thin,
                    top=med if rr == head_row else thin,
                    bottom=med if rr == bot_row else thin,
                )

    for zname, zlabel, zlocs in zones:
        r += 1   # blank spacer row before each zone band - user request
        band = sm.cell(r, 1, zlabel)
        band.font = Font(name="Calibri", size=11, bold=True, color=NAVY)
        band.alignment = Alignment(horizontal="left", vertical="center")
        sm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
        for cc in range(1, last_col + 1):
            sm.cell(r, cc).fill = PatternFill("solid", fgColor=ZONE_BAND_FILL)
        sm.row_dimensions[r].height = 15
        r += 2   # blank spacer row after the band too, before the first block - user request
        top0 = r
        gaps = LAYOUT_GAPS.get(zname, {})
        padded = []
        for loc in zlocs:
            padded.extend([None] * gaps.get(loc["key"], 0))
            padded.append(loc)
        for i, loc in enumerate(padded):
            if loc is None:
                continue
            lane = i % 3
            grp = i // 3
            top = top0 + grp * BLOCK_H
            _write_block(sm, loc, LANE_STARTS[lane], top, F, FB, thin, med)
        r = top0 + ((len(padded) + 2) // 3) * BLOCK_H

    for rr in range(1, r + 2):
        sm.row_dimensions[rr].height = 13 if sm.row_dimensions[rr].height is None else sm.row_dimensions[rr].height
    sm.sheet_view.showGridLines = False
    sm.freeze_panes = None   # unfrozen - user request
    for rr in range(1, sm.max_row + 1):   # columns A / M / Y: everything left-aligned - user request
        for cc in LANE_STARTS.values():
            cell = sm.cell(rr, cc)
            if cell.value is not None:
                cell.alignment = Alignment(horizontal="left", vertical=cell.alignment.vertical)

    _write_teu_summary(wb, model)
    if empty_raw is not None and len(empty_raw):
        _write_empty_repo(wb, model, empty_raw)
    for w in wb.worksheets:
        w.sheet_view.tabSelected = False
    wb.active = wb.worksheets.index(sm)   # open on Summary - user request
    sm.sheet_view.tabSelected = True
    wb.save(out)


SIZE20 = ["20'GP", "20'RE", "20'OT", "20'FR"]
SIZE40 = ["40'GP", "40'HC", "40'RH", "40'OT", "40'FR"]


def teu_rows(model):
    """Per-location 20'/40' empty-stock counts and TEUs (20'=1 TEU, 40'=2 TEU)."""
    rows = []
    t20 = t40 = 0
    for l in model["locations"]:
        c20 = sum(l["types"][d]["stock"] for d in SIZE20)
        c40 = sum(l["types"][d]["stock"] for d in SIZE40)
        rows.append((l["code"], c20, c40, c20 + c40, c20 + 2 * c40))
        t20 += c20
        t40 += c40
    rows.append(("TOTAL", t20, t40, t20 + t40, t20 + 2 * t40))
    return rows


def _write_teu_summary(wb, model):
    """A standing "20'/40' + TEU per location" recap sheet - user request."""
    ws = wb.create_sheet("TEU Summary")
    F = Font(name="Calibri", size=11)
    FB = Font(name="Calibri", size=11, bold=True)
    ws.append(["Location", "20' count", "40' count", "Total units", "Total TEUs"])
    for c in ws[1]:
        c.font = FB
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center")
    rows = teu_rows(model)
    for code, c20, c40, units, teu in rows:
        is_total = code == "TOTAL"
        ws.append([code, c20, c40, units, teu])
        r = ws.max_row
        for cc in range(1, 6):
            cell = ws.cell(r, cc)
            cell.font = FB if is_total else F
            cell.alignment = Alignment(horizontal="center" if cc > 1 else "left")
            if is_total:
                cell.fill = PatternFill("solid", fgColor=AV_BG)
    widths = [16, 11, 11, 13, 12]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    for r in range(1, ws.max_row + 1):
        ws.row_dimensions[r].height = 13
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"


def _write_empty_repo(wb, model, empty_raw):
    """"Empty Repo TH" sheet: inbound empties by ETA / yard / vessel, COUNTIFS over EmptyRaw."""
    F = Font(name="Calibri", size=11)
    FB = Font(name="Calibri", size=11, bold=True)
    thin = Side(style="thin", color="C9D2DF")
    er = wb.create_sheet("EmptyRaw")
    er.append(["Location", "ETA", "Vessel", "Voyage", "POL", "Type"])
    for _, rr in empty_raw.iterrows():
        er.append([rr["Location"], pd.to_datetime(rr["ETA"]).to_pydatetime(),
                   rr["Vessel"], rr["Voyage"], rr["POL"], rr["Type"]])
    for row in er.iter_rows():
        for c in row:
            c.font = F
        if row[0].row > 1:
            row[1].number_format = "yyyy-mm-dd"
    for i, w in enumerate([12, 12, 9, 9, 9, 8]):
        er.column_dimensions[get_column_letter(i + 1)].width = w

    ws = wb.create_sheet("Empty Repo TH")
    ws["A1"] = "Empty repo to TH  -  inbound empty containers by arrival date"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=15)
    for cc in range(1, 16):
        ws.cell(1, cc).fill = PatternFill("solid", fgColor=NAVY)
    heads = ["ETA", "Location", "Vessel", "Voyage", "POL"] + SIZE_CODES + ["Total"]
    for i, h in enumerate(heads):
        c = ws.cell(3, 1 + i, h)
        disp = SIZE2DISP.get(h)
        col = RE_FONT if disp in RE_DISP else OTFR_FONT if disp in OTFR_DISP else NAVY
        c.font = Font(name="Calibri", size=11, bold=True, color=col)
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=Side(style="medium", color=NAVY))
    groups = empty_repo_groups(empty_raw)
    r = 3
    for g in groups:
        r += 1
        ws.cell(r, 1, dt.datetime.strptime(g["eta"], "%Y-%m-%d")).number_format = "yyyy-mm-dd"
        ws.cell(r, 2, g["loc"])
        ws.cell(r, 3, g["vessel"])
        ws.cell(r, 4, g["voyage"])
        ws.cell(r, 5, g["pol"])
        for i, code in enumerate(SIZE_CODES):
            ws.cell(r, 6 + i, (f'=COUNTIFS(EmptyRaw!$A:$A,$B{r},EmptyRaw!$B:$B,$A{r},'
                               f'EmptyRaw!$C:$C,$C{r},EmptyRaw!$D:$D,$D{r},'
                               f'EmptyRaw!$E:$E,$E{r},EmptyRaw!$F:$F,{get_column_letter(6 + i)}$3)'))
        ws.cell(r, 15, f"=SUM(F{r}:N{r})")
        for cc in range(1, 16):
            c = ws.cell(r, cc)
            c.font = FB if cc in (2, 15) else F
            c.alignment = Alignment(horizontal="left" if cc <= 5 else "center")
            c.border = Border(bottom=thin)
    first, last = 4, r
    r += 1
    ws.cell(r, 1, "TOTAL")
    for cc in range(6, 16):
        L = get_column_letter(cc)
        ws.cell(r, cc, f"=SUM({L}{first}:{L}{last})")
    for cc in range(1, 16):
        c = ws.cell(r, cc)
        c.font = FB
        c.fill = PatternFill("solid", fgColor=AV_BG)
        c.alignment = Alignment(horizontal="left" if cc <= 5 else "center")
    for i, w in enumerate([12, 11, 9, 9, 9] + [7] * 9 + [8]):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"


def _write_block(sm, loc, c0, top, F, FB, thin, med):
    label = loc["code"]
    key = loc["key"]
    cN = c0 + 1
    cT0 = c0 + 2
    r_title, r_head, r_stock, r_pend = top, top + 1, top + 2, top + 3
    r_t0, r_rest, r_av1 = top + 4, top + 5, top + 6
    r_w2, r_av2, r_w3, r_av3, r_w4, r_av4 = [top + i for i in range(7, 13)]

    yard = YARD_NAME.get(label)
    tc = sm.cell(r_title, c0, f"{label}   {yard}" if yard else label)
    tc.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    tc.alignment = Alignment(horizontal="left")
    sm.merge_cells(start_row=r_title, start_column=c0, end_row=r_title, end_column=c0 + 10)
    for cc in range(c0, c0 + 11):
        sm.cell(r_title, cc).fill = PatternFill("solid", fgColor=NAVY)

    sm.cell(r_head, c0, "WK").font = FB
    for i, disp in enumerate(DISP_ORDER):
        hc = sm.cell(r_head, cT0 + i, disp)
        col = RE_FONT if disp in RE_DISP else OTFR_FONT if disp in OTFR_DISP else NAVY
        hc.font = Font(name="Calibri", size=11, bold=True, color=col)
        hc.alignment = Alignment(horizontal="center")

    sm.cell(r_stock, cN, "Stock empty in yard").font = FB
    for i, disp in enumerate(DISP_ORDER):
        sc = sm.cell(r_stock, cT0 + i, loc["types"][disp]["stock"])
        sc.font = F
        sc.alignment = Alignment(horizontal="center")
    for cc in range(c0, c0 + 11):
        sm.cell(r_stock, cc).fill = PatternFill("solid", fgColor=STOCK_BG)

    def sumifs(bl, extra):
        return (f'=SUMIFS(BookingRaw!${bl}$2:${bl}$1200,BookingRaw!$A$2:$A$1200,"{key}",'
                f'BookingRaw!$B$2:$B$1200,{extra})')

    sm.cell(r_pend, c0, "=Control!$A$4").font = F
    sm.cell(r_pend, cN, "Booking Pending pick up").font = F
    sm.cell(r_t0, c0, "=Control!$A$4").font = F
    sm.cell(r_t0, cN, f'="Booking on "&DAY(Control!$B$1)&" "&{mn("Control!$B$1")}').font = F
    sm.cell(r_rest, c0, "=Control!$A$4").font = F
    sm.cell(r_rest, cN,
            f'=IF(Control!$B$9=Control!$C$4,"Booking on "&DAY(Control!$B$9)&" "&{mn("Control!$B$9")},'
            f'"Booking on "&{range_label("Control!$B$9", "Control!$C$4")})').font = F
    sm.cell(r_av1, cN, f'="AV Balance till "&DAY(Control!$C$4)&" "&{mn("Control!$C$4")}').font = FB
    for rr, wkc, mon, sat in ((r_w2, "$A$5", "Control!$B$5", "Control!$C$5"),
                              (r_w3, "$A$6", "Control!$B$6", "Control!$C$6"),
                              (r_w4, "$A$7", "Control!$B$7", "Control!$C$7")):
        sm.cell(rr, c0, f"=Control!{wkc}").font = F
        sm.cell(rr, cN, f'="Booking on "&{range_label(mon, sat)}').font = F
    for rr, cref in ((r_av2, "$C$5"), (r_av3, "$C$6"), (r_av4, "$C$7")):
        sm.cell(rr, cN, f'="AV Balance till "&DAY(Control!{cref})&" "&{mn("Control!"+cref)}').font = FB

    for i, disp in enumerate(DISP_ORDER):
        cc = cT0 + i
        TL = get_column_letter(cc)
        BL = get_column_letter(3 + i)
        for rr, extra in (
            (r_pend, '"<"&Control!$B$1'),
            (r_t0, '"="&Control!$B$1'),
            (r_rest, '">="&Control!$B$9,BookingRaw!$B$2:$B$1200,"<="&Control!$C$4'),
            (r_w2, '">="&Control!$B$5,BookingRaw!$B$2:$B$1200,"<="&Control!$C$5'),
            (r_w3, '">="&Control!$B$6,BookingRaw!$B$2:$B$1200,"<="&Control!$C$6'),
            (r_w4, '">="&Control!$B$7,BookingRaw!$B$2:$B$1200,"<="&Control!$C$7'),
        ):
            e = sm.cell(rr, cc, sumifs(BL, extra))
            e.font = Font(name="Calibri", size=11,
                          color=RE_FONT if disp in RE_DISP else OTFR_FONT if disp in OTFR_DISP else "000000")
            e.alignment = Alignment(horizontal="center")
        sm.cell(r_av1, cc, f"={TL}{r_stock}-{TL}{r_pend}-{TL}{r_t0}-{TL}{r_rest}")
        sm.cell(r_av2, cc, f"={TL}{r_av1}-{TL}{r_w2}")
        sm.cell(r_av3, cc, f"={TL}{r_av2}-{TL}{r_w3}")
        sm.cell(r_av4, cc, f"={TL}{r_av3}-{TL}{r_w4}")
        avv = loc["types"][disp]["av"]
        for k, rr in enumerate((r_av1, r_av2, r_av3, r_av4)):
            cell = sm.cell(rr, cc)
            neg = avv[k] < 0
            cell.font = Font(name="Calibri", size=11, bold=neg, color=NEG_FONT if neg else "000000")
            cell.alignment = Alignment(horizontal="center")

    for rr in (r_av1, r_av2, r_av3, r_av4):
        for cc in range(c0, c0 + 11):
            sm.cell(rr, cc).fill = PatternFill("solid", fgColor=AV_BG)

    for rr in range(r_title, r_av4 + 1):
        for cc in range(c0, c0 + 11):
            cell = sm.cell(rr, cc)
            cell.border = Border(
                left=med if cc == c0 else thin,
                right=med if cc == c0 + 10 else thin,
                top=med if rr == r_title else thin,
                bottom=med if rr == r_av4 else thin,
            )


# --------------------------------------------------------------------------- #
def _md(iso):
    d = dt.date.fromisoformat(iso)
    return f"{d.day} {MONTHS[d.month - 1]}"


_HTML_CSS = """
:root{
  color-scheme:light dark;
  --bg:#F3F6FA;--surface:#FFFFFF;--surface-2:#F7F9FC;--ink:#14202E;--ink-2:#46586B;--muted:#7C8A9A;
  --line:#E1E7EF;--navy:#0C2340;--navy-2:#16365F;
  --gp:#0C2340;--re:#1F5FBF;--otfr:#8A5410;
  --bad:#C22A2A;--bad-bg:#FDECEC;--bad-line:#F3C2C2;--ok:#1E8E5A;--ok-bg:#E7F5EE;
  --hi:#FFF6C8;--shadow:0 1px 2px rgba(12,35,64,.06),0 4px 14px rgba(12,35,64,.05);
}
@media (prefers-color-scheme:dark){:root{
  --bg:#0D131A;--surface:#161F2A;--surface-2:#1B2633;--ink:#E9EFF6;--ink-2:#AAB8C7;--muted:#7F8E9F;
  --line:#26333F;--navy:#0A1726;--navy-2:#12263F;
  --gp:#C9D6E6;--re:#7DB0FF;--otfr:#E0AE66;
  --bad:#FF7B7B;--bad-bg:#3A1D21;--bad-line:#6B2F35;--ok:#4CC38A;--ok-bg:#12301F;
  --hi:#3B3516;--shadow:0 1px 2px rgba(0,0,0,.4);
}}
*{box-sizing:border-box}
html{scroll-behavior:smooth}
body{margin:0;font:14px/1.5 "Segoe UI","Leelawadee UI",Tahoma,system-ui,sans-serif;color:var(--ink);background:var(--bg);
  -webkit-font-smoothing:antialiased}
.wrap{max-width:1280px;margin:0 auto;padding:0 20px}
header.top{background:linear-gradient(135deg,var(--navy) 0%,var(--navy-2) 100%);color:#fff;padding:22px 0 26px}
header.top h1{margin:0;font-size:22px;font-weight:700;letter-spacing:.2px}
header.top .sub{margin-top:4px;color:#B9C8DB;font-size:13px}
header.top .date{display:inline-block;margin-top:12px;background:rgba(255,255,255,.12);border:1px solid rgba(255,255,255,.22);
  border-radius:999px;padding:3px 12px;font-size:13px;font-weight:600}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));gap:14px;margin-top:-18px}
.kpi{background:var(--surface);border:1px solid var(--line);border-radius:12px;padding:14px 16px;box-shadow:var(--shadow)}
.kpi .lbl{font-size:12px;color:var(--ink-2);font-weight:600}
.kpi .val{font-size:30px;line-height:1.15;font-weight:700;margin-top:2px;font-variant-numeric:tabular-nums}
.kpi .hint{font-size:11.5px;color:var(--muted);margin-top:2px}
.kpi.bad .val{color:var(--bad)}.kpi.ok .val{color:var(--ok)}
section.panel{margin-top:22px}
section.panel>h2{font-size:16px;margin:0 0 4px;display:flex;align-items:center;gap:8px}
section.panel>.lead{margin:0 0 10px;color:var(--ink-2);font-size:12.5px}
.box{background:var(--surface);border:1px solid var(--line);border-radius:12px;box-shadow:var(--shadow);overflow:hidden}
.alerts{width:100%;border-collapse:collapse}
.alerts th{font-size:11.5px;color:var(--muted);font-weight:600;text-align:left;padding:9px 14px;background:var(--surface-2);border-bottom:1px solid var(--line)}
.alerts td{padding:9px 14px;border-bottom:1px solid var(--line);vertical-align:middle}
.alerts tr:last-child td{border-bottom:0}
.alerts tr.alert-row{cursor:pointer}.alerts tr.alert-row:hover td{background:var(--surface-2)}
.alerts .loc{font-weight:700;white-space:nowrap}
.alerts .loc a{color:inherit;text-decoration:none}
.alerts .when b{font-weight:600}.alerts .when small{color:var(--muted);margin-left:6px}
.gap{display:flex;align-items:center;gap:10px;min-width:170px}
.gap .bar{flex:1;height:8px;border-radius:4px;background:var(--bad-bg);overflow:hidden;min-width:60px}
.gap .bar i{display:block;height:100%;background:var(--bad);border-radius:4px}
.gap b{color:var(--bad);font-variant-numeric:tabular-nums;white-space:nowrap}
.allgood{padding:16px;display:flex;gap:10px;align-items:center;color:var(--ok);font-weight:600}
details.more summary{cursor:pointer;padding:10px 14px;font-size:12.5px;color:var(--ink-2);border-top:1px solid var(--line);background:var(--surface-2)}
.controls{position:sticky;top:0;z-index:5;background:color-mix(in srgb,var(--bg) 92%,transparent);backdrop-filter:blur(6px);
  padding:10px 0;margin-top:14px;display:flex;gap:8px;flex-wrap:wrap;align-items:center;border-bottom:1px solid var(--line)}
.controls .lab{font-size:12px;color:var(--muted);margin:0 2px 0 6px}
.controls button{border:1px solid var(--line);background:var(--surface);color:var(--ink);border-radius:999px;padding:6px 14px;cursor:pointer;font-size:13px;font-weight:600}
.controls button:hover{border-color:var(--navy-2)}
.controls button.on{background:var(--navy);color:#fff;border-color:var(--navy)}
@media (prefers-color-scheme:dark){.controls button.on{background:#2A5CA8;border-color:#2A5CA8}}
.zone-h{display:flex;align-items:baseline;gap:10px;margin:26px 0 10px}
.zone-h h2{margin:0;font-size:18px}.zone-h small{color:var(--muted);font-size:13px}
.zone-h .tag{font-size:11px;font-weight:700;padding:2px 9px;border-radius:999px;background:var(--hi);color:#5B4A00}
@media (prefers-color-scheme:dark){.zone-h .tag{color:#F2DD87}}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(470px,1fr));gap:16px}
.card{background:var(--surface);border:1px solid var(--line);border-radius:12px;box-shadow:var(--shadow);overflow:hidden;scroll-margin-top:70px}
.card.short{border-color:var(--bad-line)}
.chd{display:flex;justify-content:space-between;align-items:flex-start;gap:10px;padding:12px 14px;border-bottom:1px solid var(--line);background:var(--surface-2)}
.card.short .chd{background:var(--bad-bg)}
.chd .code{font-size:17px;font-weight:700}
.chd .yard{font-size:12px;color:var(--ink-2)}
.chd .meta{font-size:12px;color:var(--muted);margin-top:2px}
.pill{font-size:11.5px;font-weight:700;padding:3px 10px;border-radius:999px;white-space:nowrap}
.pill.bad{background:var(--bad);color:#fff}.pill.ok{background:var(--ok-bg);color:var(--ok)}
@media (prefers-color-scheme:dark){.pill.bad{color:#1A0B0D}}
.scroll{overflow-x:auto}
table.t{width:100%;border-collapse:collapse;font-size:13px}
table.t th,table.t td{padding:7px 8px;text-align:right;border-bottom:1px solid var(--line);white-space:nowrap}
table.t thead th{font-size:11.5px;font-weight:600;color:var(--muted);vertical-align:bottom;background:var(--surface)}
table.t thead th small{display:block;font-weight:500;font-size:10.5px}
table.t thead th.cur{color:var(--ink)}
table.t tbody tr:last-child th,table.t tbody tr:last-child td{border-bottom:0}
table.t th.type,table.t td.type{text-align:left;font-weight:700;padding-left:14px}
table.t th.type small{font-weight:500;color:var(--muted);margin-left:6px;font-size:11px}
.dot{display:inline-block;width:9px;height:9px;border-radius:50%;margin-right:7px;background:var(--gp)}
.re .dot{background:var(--re)}.otfr .dot{background:var(--otfr)}
td.num{font-variant-numeric:tabular-nums}td.num.z{color:var(--muted)}
td.stock b{font-size:14px}
td.wk{font-variant-numeric:tabular-nums;line-height:1.2}
td.wk b{display:block;font-size:14px}
td.wk small{display:block;color:var(--muted);font-size:10.5px}
td.wk.neg{background:var(--bad-bg);box-shadow:inset 3px 0 0 var(--bad)}
td.wk.neg b{color:var(--bad)}td.wk.neg b::before{content:"▼ ";font-size:9px;vertical-align:1px}
td.trend{padding-right:14px;text-align:center}
svg.sp{display:block;margin:auto}
.nobk{width:100%;border-collapse:collapse;font-size:13px}
.nobk th,.nobk td{padding:8px 10px;text-align:center;border-bottom:1px solid var(--line);white-space:nowrap}
.nobk thead th{background:var(--surface-2);font-size:12px;color:var(--ink-2)}
.nobk thead th.re{color:var(--re)}.nobk thead th.otfr{color:var(--otfr)}
.nobk tbody th{text-align:left;font-weight:700}
.nobk td.z{color:var(--muted)}.nobk td.tot,.nobk th.tot{font-weight:700;background:var(--surface-2)}
.nobk tr.sum td,.nobk tr.sum th{background:var(--hi);font-weight:700;border-bottom:0}
.nobk td.eta b{display:block}.nobk td.eta small{color:var(--muted);font-size:11px}
.nobk .l{text-align:left}
.help{margin-top:22px}
.help summary{cursor:pointer;font-weight:700;font-size:14px;padding:12px 16px}
.help .body{padding:0 16px 14px;color:var(--ink-2);font-size:13px}
.help ul{margin:6px 0 0;padding-left:18px}.help li{margin:3px 0}
footer{padding:22px 0 40px;color:var(--muted);font-size:12px}
.hide{display:none!important}
@media (max-width:620px){
  .wrap{padding:0 12px}.grid{grid-template-columns:1fr}header.top h1{font-size:19px}
  .kpi .val{font-size:26px}.alerts th,.alerts td{padding:8px 9px}.gap{min-width:110px}
}
"""

_HTML_JS = """
var st={zone:'all',short:false};
function apply(){
  document.querySelectorAll('.card').forEach(function(c){
    var ok=(st.zone==='all'||c.dataset.zone===st.zone)&&(!st.short||c.dataset.badge==='short');
    c.classList.toggle('hide',!ok);
  });
  document.querySelectorAll('.zone-sec[data-zone]').forEach(function(z){
    var any=z.querySelector('.card:not(.hide)');
    z.classList.toggle('hide',!any);
  });
  document.querySelectorAll('.alert-row').forEach(function(r){
    r.classList.toggle('hide',st.zone!=='all'&&r.dataset.zone!==st.zone);
  });
}
document.querySelectorAll('.controls button').forEach(function(b){
  b.onclick=function(){
    if(b.dataset.f==='zone'){st.zone=b.dataset.v;
      document.querySelectorAll('.controls button[data-f="zone"]').forEach(function(x){x.classList.toggle('on',x===b)});}
    else{st.short=!st.short;b.classList.toggle('on',st.short);}
    apply();
  };
});
document.querySelectorAll('.alert-row').forEach(function(r){
  r.onclick=function(){
    var c=document.getElementById(r.dataset.target);
    if(c){c.classList.remove('hide');c.scrollIntoView({behavior:'smooth',block:'start'});}
  };
});
"""


def _fmt(n):
    return f"{n:,}".replace("-", "−")


def _rng(a, b):
    da, db = dt.date.fromisoformat(a), dt.date.fromisoformat(b)
    if da.month == db.month:
        return f"{da.day}–{db.day} {MONTHS[db.month - 1]}"
    return f"{da.day} {MONTHS[da.month - 1]}–{db.day} {MONTHS[db.month - 1]}"


def _spark(vals):
    w, h, pad = 64, 22, 3
    lo, hi = min(0, min(vals)), max(max(vals), 1)
    span = (hi - lo) or 1
    xs = [pad + i * (w - 2 * pad) / (len(vals) - 1) for i in range(len(vals))]
    ys = [pad + (hi - v) * (h - 2 * pad) / span for v in vals]
    pts = " ".join(f"{x:.1f},{y:.1f}" for x, y in zip(xs, ys))
    bad = min(vals) < 0
    col = "var(--bad)" if bad else "var(--ink-2)"
    zero = ""
    if lo < 0:
        zy = pad + hi * (h - 2 * pad) / span
        zero = (f"<line x1='0' x2='{w}' y1='{zy:.1f}' y2='{zy:.1f}' stroke='var(--muted)' "
                "stroke-width='1' stroke-dasharray='2 2'/>")
    return (f"<svg class='sp' width='{w}' height='{h}' viewBox='0 0 {w} {h}' aria-hidden='true'>{zero}"
            f"<polyline points='{pts}' fill='none' stroke='{col}' stroke-width='2' stroke-linejoin='round' "
            f"stroke-linecap='round'/><circle cx='{xs[-1]:.1f}' cy='{ys[-1]:.1f}' r='2.6' fill='{col}'/></svg>")


def build_html(model, out: Path):
    import re
    wk = model["iso_weeks"]
    weeks = model["weeks"]
    today = dt.date.fromisoformat(model["date"])
    rdate = _md(model["date"]) + " " + model["date"][:4]
    zones = [("BKK", "Bangkok area depots"), ("LCH", "Laem Chabang area depots")]
    E = html.escape

    def cls(disp):
        return "re" if disp in RE_DISP else "otfr" if disp in OTFR_DISP else "gp"

    def anchor(l):
        return "loc-" + re.sub(r"\W+", "-", l["code"]).strip("-")

    with_bkg = [l for l in model["locations"]
                if l.get("has_booking", True) or l["key"] in ALWAYS_FULL_BLOCK]
    no_bkg = [l for l in model["locations"]
              if not l.get("has_booking", True) and l["key"] not in ALWAYS_FULL_BLOCK]
    all_t = [t for l in model["locations"] for t in l["types"].values()]
    total_stock = sum(t["stock"] for t in all_t)
    total_pend = sum(t["pending"] for t in all_t)

    # ---- shortage alerts ------------------------------------------------
    alerts = []
    for l in with_bkg:
        for d in DISP_ORDER:
            av = l["types"][d]["av"]
            if min(av) < 0:
                first = next(i for i, v in enumerate(av) if v < 0)
                alerts.append((min(av), l, d, first))
    alerts.sort(key=lambda a: a[0])
    short_locs = {a[1]["code"] for a in alerts}
    worst_all = abs(alerts[0][0]) if alerts else 1

    def alert_row(a):
        mn, l, d, first = a
        when = "this week" if first == 0 else f"from WK{wk[first]}"
        pct = max(6, round(abs(mn) / worst_all * 100))
        return (f"<tr class='alert-row' data-zone='{l['zone']}' data-target='{anchor(l)}'>"
                f"<td class='loc'><a href='#{anchor(l)}'>{E(l['code'])}</a></td>"
                f"<td class='{cls(d)}'><i class='dot' style='display:inline-block'></i>"
                f"<b>{E(d)}</b> <small style='color:var(--muted)'>{DISP2SIZE[d]}</small></td>"
                f"<td class='when'><b>{when}</b><small>by {_md(weeks[first][1])}</small></td>"
                f"<td><div class='gap'><div class='bar'><i style='width:{pct}%'></i></div>"
                f"<b>{_fmt(mn)}</b></div></td></tr>")

    if alerts:
        rows = [alert_row(a) for a in alerts]
        head = ("<tr><th>Location</th><th>Container type</th><th>Runs short</th>"
                "<th>Lowest balance (shortfall)</th></tr>")
        tail = ""
        if len(rows) > 10:
            tail = (f"<details class='more'><summary>Show {len(rows) - 10} more</summary>"
                    f"<table class='alerts'><tbody>{''.join(rows[10:])}</tbody></table></details>")
        alert_html = (f"<div class='box'><table class='alerts'><thead>{head}</thead>"
                      f"<tbody>{''.join(rows[:10])}</tbody></table>{tail}</div>")
        lead = ("ตู้ว่างไม่พอรับ booking ภายใน 4 สัปดาห์ — เรียงจากขาดมากสุด "
                "(คลิกแถวเพื่อไปที่การ์ดของลานนั้น)")
    else:
        alert_html = "<div class='box'><div class='allgood'>✔ No shortage in the next 4 weeks</div></div>"
        lead = "ไม่มีลาน/ชนิดตู้ไหนขาดใน 4 สัปดาห์ข้างหน้า"

    # ---- location cards -------------------------------------------------
    def card(l):
        t = l["types"]
        active = [d for d in DISP_ORDER if t[d]["stock"] or any(t[d]["wk"]) or t[d]["pending"]] or ["20'GP"]
        short_types = [d for d in DISP_ORDER if min(t[d]["av"]) < 0]
        badge = "short" if short_types else "ok"
        pill = (f"<span class='pill bad'>▼ SHORT · {len(short_types)} type{'s' if len(short_types) > 1 else ''}</span>"
                if short_types else "<span class='pill ok'>✔ OK</span>")
        yard = YARD_NAME.get(l["code"])
        stock_sum = sum(t[d]["stock"] for d in DISP_ORDER)
        pend_sum = sum(t[d]["pending"] for d in DISP_ORDER)
        wk_head = "".join(
            f"<th class='{'cur' if i == 0 else ''}'>WK{wk[i]}<small>{_rng(*weeks[i])}</small></th>"
            for i in range(4))
        body = []
        for d in active:
            x = t[d]
            cells = []
            for i in range(4):
                neg = x["av"][i] < 0
                bk = x["wk"][i]
                tip = (f"WK{wk[i]} {_rng(*weeks[i])}: booked {bk}, balance after = {x['av'][i]}")
                cells.append(f"<td class='wk{' neg' if neg else ''}' title='{E(tip)}'><b>{_fmt(x['av'][i])}</b>"
                             f"<small>{'−' + str(bk) if bk else '–'}</small></td>")
            spark = _spark([x["stock"] - x["pending"]] + x["av"])
            zc = " z" if not x["pending"] else ""
            body.append(
                f"<tr class='{cls(d)}'><th class='type'><i class='dot'></i>{E(d)}<small>{DISP2SIZE[d]}</small></th>"
                f"<td class='num stock'><b>{_fmt(x['stock'])}</b></td>"
                f"<td class='num{zc}'>{_fmt(x['pending']) if x['pending'] else '–'}</td>"
                f"{''.join(cells)}<td class='trend'>{spark}</td></tr>")
        yard_html = f"<div class='yard'>{E(yard)}</div>" if yard else ""
        return (f"<div class='card{' short' if short_types else ''}' id='{anchor(l)}' data-zone='{l['zone']}' "
                f"data-badge='{badge}'><div class='chd'><div><div class='code'>{E(l['code'])}</div>{yard_html}"
                f"<div class='meta'>{_fmt(stock_sum)} empty on hand · {_fmt(pend_sum)} overdue pickup</div></div>{pill}</div>"
                f"<div class='scroll'><table class='t'><thead><tr><th class='type'>Type</th><th>Stock</th>"
                f"<th>Overdue<small>pickup</small></th>{wk_head}<th>Trend</th></tr></thead>"
                f"<tbody>{''.join(body)}</tbody></table></div></div>")

    sections = []
    for zid, zname in zones:
        zlocs = [l for l in with_bkg if l["zone"] == zid]
        if not zlocs:
            continue
        n_short = sum(1 for l in zlocs if l["code"] in short_locs)
        tag = f"<span class='tag'>{n_short} at risk</span>" if n_short else ""
        sections.append(
            f"<div class='zone-sec' data-zone='{zid}'><div class='zone-h'><h2>{zid}</h2>"
            f"<small>{zname}</small>{tag}</div><div class='grid'>{''.join(card(l) for l in zlocs)}</div></div>")

    # ---- stock-on-hand-only table --------------------------------------
    if no_bkg:
        th = "".join(f"<th class='{cls(SIZE2DISP[c])}'>{c}</th>" for c in SIZE_CODES)
        trs = []
        for l in no_bkg:
            t = l["types"]
            tds = "".join(
                f"<td class='z'>·</td>" if not t[d]["stock"] else f"<td>{_fmt(t[d]['stock'])}</td>"
                for d in DISP_ORDER)
            trs.append(f"<tr><th>{E(l['code'])}</th><td class='l'>{E(l['zone'])}</td>{tds}"
                       f"<td class='tot'>{_fmt(sum(t[d]['stock'] for d in DISP_ORDER))}</td></tr>")
        sections.insert(0, (
            "<section class='panel'><h2>Stock on hand only <small style='color:var(--muted);font-weight:500'>"
            "no active bookings this period</small></h2>"
            "<p class='lead'>ลานที่มีตู้ว่างแต่ไม่มี booking ในช่วงนี้</p>"
            f"<div class='box scroll'><table class='nobk'><thead><tr><th class='l'>Location</th><th class='l'>Zone</th>"
            f"{th}<th>Total</th></tr></thead><tbody>{''.join(trs)}</tbody></table></div></section>"))

    # ---- empty repo to TH ----------------------------------------------
    er = model.get("empty_repo") or []
    er_total = sum(g["total"] for g in er)
    if er:
        th = "".join(f"<th class='{cls(SIZE2DISP[c])}'>{c}</th>" for c in SIZE_CODES)
        trs = []
        for g in er:
            days = (dt.date.fromisoformat(g["eta"]) - today).days
            when = ("today" if days == 0 else "tomorrow" if days == 1 else
                    f"in {days} days" if days > 1 else f"{-days} days ago")
            tds = "".join(f"<td class='z'>·</td>" if not g["types"][c] else f"<td>{g['types'][c]}</td>"
                          for c in SIZE_CODES)
            trs.append(f"<tr><td class='eta l'><b>{_md(g['eta'])}</b><small>{when}</small></td>"
                       f"<th>{E(g['loc'])}</th><td class='l'>{E(g['vessel'])} {E(g['voyage'])}</td>"
                       f"<td class='l'>{E(g['pol'])}</td>{tds}<td class='tot'>{g['total']}</td></tr>")
        ttl = {c: sum(g["types"][c] for g in er) for c in SIZE_CODES}
        trs.append("<tr class='sum'><th>TOTAL</th><th></th><td></td><td></td>"
                   + "".join(f"<td>{ttl[c] or '·'}</td>" for c in SIZE_CODES) + f"<td>{er_total}</td></tr>")
        etas = sorted({g["eta"] for g in er})
        span = _md(etas[0]) if len(etas) == 1 else f"{_md(etas[0])}–{_md(etas[-1])}"
        sections.insert(0, (
            f"<section class='panel'><h2>Empty repo to TH <small style='color:var(--muted);font-weight:500'>"
            f"{er_total} empty containers arriving {span}</small></h2>"
            "<p class='lead'>ตู้เปล่าที่เรือกำลังนำเข้าไทย (ยังไม่รวมในตัวเลข stock และ AV balance)</p>"
            f"<div class='box scroll'><table class='nobk'><thead><tr><th class='l'>ETA</th><th class='l'>Location</th>"
            f"<th class='l'>Vessel / Voy</th><th class='l'>POL</th>{th}<th>Total</th></tr></thead>"
            f"<tbody>{''.join(trs)}</tbody></table></div></section>"))

    # ---- KPIs -----------------------------------------------------------
    n_full = len(with_bkg)
    kpis = [
        ("", "Empty stock on hand", _fmt(total_stock), "all container types, all yards"),
        ("", "Overdue pickup", _fmt(total_pend), "bookings dated before today, not yet picked up"),
        ("bad" if short_locs else "ok", "Yards at risk", f"{len(short_locs)} / {n_full}",
         f"{len(alerts)} type(s) run short within 4 weeks" if alerts else "no shortage in 4 weeks"),
    ]
    if er:
        kpis.append(("", "Empty repo inbound", _fmt(er_total), f"arriving {span}"))
    kpi_html = "".join(
        f"<div class='kpi {c}'><div class='lbl'>{l}</div><div class='val'>{v}</div><div class='hint'>{h}</div></div>"
        for c, l, v, h in kpis)

    help_html = (
        "<details class='box help'><summary>วิธีอ่านรายงาน · How to read</summary><div class='body'><ul>"
        "<li><b>Stock</b> = ตู้เปล่าที่อยู่ในลานตอนนี้ · <b>Overdue pickup</b> = booking ที่เลยวันนัดแล้วแต่ยังไม่ได้รับตู้ (หักออกก่อน)</li>"
        "<li>ช่อง <b>WK</b>: ตัวเลขใหญ่ = ตู้คงเหลือ (AV balance) หลังหัก booking ของสัปดาห์นั้นและสัปดาห์ก่อนหน้าทั้งหมด · "
        "ตัวเล็ก = จำนวน booking ของสัปดาห์นั้น</li>"
        "<li><span style='color:var(--bad);font-weight:700'>▼ สีแดง</span> = ตู้คงเหลือติดลบ คือตู้ไม่พอรับ booking</li>"
        "<li><b>Trend</b> = เส้นแนวโน้มของตู้คงเหลือ จากวันนี้ถึงปลาย 4 สัปดาห์ (เส้นประ = ศูนย์)</li>"
        "<li>WK1 รวม booking ตั้งแต่วันนี้ถึงวันเสาร์ของสัปดาห์นี้</li></ul></div></details>")

    css = _HTML_CSS
    page = (
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>"
        "<meta name='viewport' content='width=device-width,initial-scale=1'>"
        f"<title>Daily Container Stock vs Booking · {rdate}</title><style>{css}</style></head><body>"
        f"<header class='top'><div class='wrap'><h1>Daily Container Stock vs Booking</h1>"
        f"<div class='sub'>ตู้เปล่าในลาน เทียบกับ booking 4 สัปดาห์ · all 9 container types</div>"
        f"<span class='date'>Report date · {rdate}</span></div></header>"
        f"<main class='wrap'><div class='kpis'>{kpi_html}</div>"
        f"<section class='panel'><h2>{'⚠' if alerts else '✔'} Shortage alerts</h2><p class='lead'>{lead}</p>{alert_html}</section>"
        "<div class='controls'><span class='lab'>Zone</span>"
        "<button data-f='zone' data-v='all' class='on'>All</button>"
        "<button data-f='zone' data-v='BKK'>BKK</button><button data-f='zone' data-v='LCH'>LCH</button>"
        "<span class='lab' style='margin-left:12px'>Show</span><button data-f='short'>Shortage only</button></div>"
        f"{''.join(sections)}{help_html}"
        f"<footer>Stock = empty containers in yard now (snapshot). Booking = pickups by TRAN DT. "
        f"AV balance = stock − overdue pickup − cumulative bookings through each week. "
        f"Generated from {E(model['date'])} STAYING + booking exports.</footer></main>"
        f"<script>{_HTML_JS}</script></body></html>")
    out.write_text(page, encoding="utf-8")


def find_default_inputs():
    """Auto-pick the newest STAYING file, and ALL booking-like files, from input/.

    Booking filenames vary day to day (BKG+PD, PD+BKG-3WK, PD+BKG-3WKS, ...)
    and some days the source system splits the export into a "PENDING" file
    (overdue bookings) plus a separate forward-looking "BKG-3WK" file - both
    get matched and combined (see load_all_bookings). Match loosely on "BKG"
    or "PENDING" in the name. "-RH" files are the separate reefer-only
    report's inputs and are always excluded here. Keep input/ holding just
    the current day's files (archive old ones) so this stays unambiguous.
    """
    if not INPUT_DIR.is_dir():
        return None, []

    def find(match):
        return [p for p in INPUT_DIR.glob("*.xls*")
                if match(p.stem.upper()) and "-RH" not in p.stem.upper()]

    staying = find(lambda n: "STAYING" in n)
    stock_path = max(staying, key=lambda p: p.stat().st_mtime) if staying else None
    bkg_paths = sorted(find(lambda n: "BKG" in n or "PENDING" in n))
    return stock_path, bkg_paths


def find_empty_input():
    """Newest *EMPTY* export in input/ (inbound empty repo list), or None."""
    if not INPUT_DIR.is_dir():
        return None
    cands = [p for p in INPUT_DIR.glob("*.xls*")
             if "EMPTY" in p.stem.upper() and "-RH" not in p.stem.upper()]
    return max(cands, key=lambda p: p.stat().st_mtime) if cands else None


def _is_date_str(s: str) -> bool:
    try:
        dt.date.fromisoformat(s)
        return True
    except ValueError:
        return False


# --------------------------------------------------------------------------- #
def main():
    args = sys.argv[1:]
    override_arg = args[-1] if args and _is_date_str(args[-1]) else None
    file_args = args[:-1] if override_arg else args

    if file_args:
        stock_path = FOLDER / file_args[0]
        bkg_paths = [FOLDER / a for a in file_args[1:]]
    else:
        # no explicit file paths (only maybe a trailing date) - auto-discover
        # the newest STAYING file and every booking-like file in input/.
        stock_path, bkg_paths = find_default_inputs()
        if not stock_path or not bkg_paths:
            raise SystemExit(
                f"No STAYING + booking .xls files found in {INPUT_DIR} and no explicit "
                "file paths were given. Usage:\n"
                '  python build_stock_daily.py "<stock>.xls" "<booking1>.xls" ["<booking2>.xls" ...] [YYYY-MM-DD]\n'
                "  python build_stock_daily.py [YYYY-MM-DD]   (auto from input/)")
        print(f"auto-detected input: {stock_path.name}  +  {', '.join(p.name for p in bkg_paths)}")

    # optional date override, for when the report must be labeled for a date
    # other than the machine's live clock (in that case Control!B1 is written
    # as a literal date, not a live =TODAY()).
    override = dt.date.fromisoformat(override_arg) if override_arg else None
    today = override or dt.date.today()
    tag = f"{today.month}-{today.day}"

    stock_agg, stock_raw = load_stock(stock_path)
    bkg = load_all_bookings(bkg_paths)
    model = build_model(stock_agg, bkg, today)
    empty_path = find_empty_input()
    empty_raw = load_empty_repo(empty_path) if empty_path else None
    if empty_raw is not None:
        model["empty_repo"] = empty_repo_groups(empty_raw)
        print(f"empty repo: {empty_path.name}  ({len(empty_raw)} containers)")

    (FOLDER / f"Stock_Daily_{tag}.model.json").write_text(
        json.dumps(model, indent=1, ensure_ascii=False), encoding="utf-8")
    build_excel(model, stock_agg, stock_raw, bkg, FOLDER / f"Stock_Daily_{tag}.xlsx",
                override_date=override, empty_raw=empty_raw)
    print("xlsx:", f"Stock_Daily_{tag}.xlsx")
    build_html(model, FOLDER / f"Stock_Daily_{tag}.html")
    print("html:", f"Stock_Daily_{tag}.html")


if __name__ == "__main__":
    main()
